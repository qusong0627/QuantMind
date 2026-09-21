"""回测报告导出件必须带免责段（PDF 与 Excel 两条链路）。

缺口（2026-09-21 核实）：``PDFReportGenerator`` / ``ExcelReportGenerator``
此前**没有任何免责段** —— 导出的是一份「没有出处的收益曲线与指标表」。
这两个文件会被下载、转发、打印，屏幕上那条免责横幅跟不出去。

**为什么跑完整个生成器再从产物里读回来**：只测 ``disclaimer_rows`` 纯函数会让
「忘了接线」这类缺陷整类溜过（那类用例在 `test_export_disclaimer.py` 里已有）。
这里要证明的是接线本身 —— 生成器真的把免责段放进了交付物。
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from reportlab.platypus import SimpleDocTemplate

from backend.services.engine.qlib_app.services.report_generator import (
    ExcelReportGenerator,
    PDFReportGenerator,
    _result_data_range,
)
from backend.shared.export_disclaimer import DISCLAIMER_LABELS, DISCLAIMER_SENTENCE

#: 同时带顶层与 config 两套 start/end —— 两种结果形态都在用（见 _result_data_range）
RESULT: dict = {
    "start_date": "2026-01-01",
    "end_date": "2026-09-18",
    "config": {"start_date": "2026-01-01", "end_date": "2026-09-18"},
}

SHEET_NAME = "免责声明"


def _disclaimer_cells(xlsx_bytes: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert SHEET_NAME in wb.sheetnames, (
        f"Excel 报告缺「{SHEET_NAME}」表: {wb.sheetnames}"
    )
    return [
        [c for c in row if c is not None]
        for row in wb[SHEET_NAME].iter_rows(values_only=True)
        if any(c is not None for c in row)
    ]


# ── Excel：产物里真有这张表 ──────────────────────────────────────────


def test_excel_report_carries_disclaimer_sheet() -> None:
    xlsx = ExcelReportGenerator().generate(RESULT)
    cells = _disclaimer_cells(xlsx)

    flat = [" ".join(str(c) for c in row) for row in cells]
    assert any(DISCLAIMER_SENTENCE in line for line in flat), f"缺免责语句: {flat}"

    labels = [row[0] for row in cells]
    assert DISCLAIMER_LABELS["exportedAt"] in labels, f"缺生成时间行: {labels}"
    assert DISCLAIMER_LABELS["dataRange"] in labels, f"缺数据区间行: {labels}"

    range_row = next(row for row in cells if row[0] == DISCLAIMER_LABELS["dataRange"])
    assert range_row[1] == "2026-01-01 ~ 2026-09-18"


def test_excel_report_omits_range_when_result_has_none() -> None:
    """结果里没有起止日期时**不出现**区间行 —— 不写空值、不写半截区间。"""
    cells = _disclaimer_cells(ExcelReportGenerator().generate({"config": {}}))
    labels = [row[0] for row in cells]
    assert DISCLAIMER_LABELS["dataRange"] not in labels
    # 但生成时间与语句照旧在（不是"整段没写"）
    assert DISCLAIMER_LABELS["exportedAt"] in labels
    assert any(DISCLAIMER_SENTENCE in str(row) for row in cells)


def test_excel_report_still_has_data_sheets() -> None:
    """免责段是**加**一张表，不是挤掉原来的表（防"加了免责、丢了数据"）。"""
    wb = load_workbook(io.BytesIO(ExcelReportGenerator().generate(RESULT)))
    assert wb.sheetnames[:5] == [
        "核心指标",
        "权益曲线",
        "交易明细",
        "持仓明细",
        "每日收益率",
    ]
    assert wb.sheetnames[-1] == SHEET_NAME


# ── PDF：交给 doc.build 的 story 里真有免责段 ────────────────────────


def test_pdf_generate_puts_disclaimer_into_story(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """拦下 `SimpleDocTemplate.build` 看 story —— PDF 字节是压缩流，grep 不出文字，
    所以断言「交给渲染器的内容」而不是「渲染后的字节」。

    这条能抓住的缺陷：免责段在辅助函数里写对了、但 `generate()` 忘了 append。
    """
    captured: list = []

    def _capture(self, story, *args, **kwargs):  # noqa: ANN001, ARG001
        captured.extend(story)

    monkeypatch.setattr(SimpleDocTemplate, "build", _capture)
    PDFReportGenerator().generate(RESULT)

    assert captured, "story 为空 —— 本用例失去意义（跑完整个 generate 才有意义）"
    texts = [getattr(flowable, "text", "") or "" for flowable in captured]
    assert any(DISCLAIMER_SENTENCE in t for t in texts), "story 里没有免责语句"
    assert any("2026-01-01 ~ 2026-09-18" in t for t in texts), "story 里没有数据区间"
    assert any(DISCLAIMER_LABELS["exportedAt"] in t for t in texts), (
        "story 里没有生成时间"
    )


# ── 区间取数：两种结果形态都要认 ─────────────────────────────────────


@pytest.mark.parametrize(
    ("result", "expected"),
    [
        ({"start_date": "a", "end_date": "b"}, ("a", "b")),
        ({"config": {"start_date": "a", "end_date": "b"}}, ("a", "b")),
        (
            {"start_date": "a", "config": {"start_date": "z", "end_date": "b"}},
            ("a", "b"),
        ),
        ({"config": None, "start_date": "a"}, ("a", None)),
        ({}, (None, None)),
    ],
)
def test_result_data_range_covers_both_shapes(result: dict, expected: tuple) -> None:
    """顶层优先、config 兜底 —— 两种形态都在用，取值规则只写一处。"""
    assert _result_data_range(result) == expected
